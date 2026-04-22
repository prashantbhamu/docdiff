import sys
import os
import re
import difflib
import unicodedata
import tkinter as tk
from tkinter import filedialog

# ---------------------------------------------------------------------------
# Prerequisites:
# You will need to install two libraries to extract text from PDFs and Word docs.
# Run this in your terminal before executing the script:
# pip install PyPDF2 python-docx
# ---------------------------------------------------------------------------

try:
    import PyPDF2
except ImportError:
    print("Error: PyPDF2 is not installed.")
    print("Please install it by running: pip install PyPDF2")
    sys.exit(1)

try:
    import docx
except ImportError:
    print("Error: python-docx is not installed.")
    print("Please install it by running: pip install python-docx")
    sys.exit(1)


def sanitize_text(text):
    """
    Aggressively cleans text of invisible characters, decomposes ligatures, 
    and normalizes unicode punctuation and spacing that commonly cause false positive differences.
    """
    if not text:
        return ""
    
    # 1. Normalize unicode (fixes ligatures like 'fi', 'fl' and standardizes characters)
    text = unicodedata.normalize('NFKC', text)
    
    # 2. Strip ALL control and format characters (Categories starting with 'C')
    # This safely catches \u00AD (soft hyphen), \u200B (zero-width space), LTR/RTL marks, and PDF hidden bytes.
    text = "".join(ch for ch in text if unicodedata.category(ch)[0] != "C")
    
    # 3. Normalize various dashes/hyphens to standard ASCII hyphen
    text = re.sub(r'[\u2010\u2011\u2012\u2013\u2014\u2015\u2212]', '-', text)
    
    # 4. Normalize smart quotes to standard straight quotes
    text = re.sub(r'[\u2018\u2019\u201A\u201B]', "'", text)
    text = re.sub(r'[\u201C\u201D\u201E\u201F]', '"', text)
    
    # 5. Fix common PDF extraction spacing bugs around punctuation
    # Removes erratic spaces before closing punctuation (e.g., "word ." -> "word.")
    text = re.sub(r'\s+([.,;:!?\]\)])', r'\1', text) 
    # Removes erratic spaces after opening brackets (e.g., "( word" -> "(word")
    text = re.sub(r'([\[\(])\s+', r'\1', text)
    
    # 6. Collapse all remaining whitespace into a single space
    text = re.sub(r'\s+', ' ', text).strip()
    
    return text


def is_false_positive(s1, s2):
    """
    Checks if the difference between two strings is purely due to 
    spacing, capitalization, hyphens, or punctuation (e.g. "pow er" vs "power").
    It strips all non-alphanumeric characters to compare the true core text.
    """
    norm1 = re.sub(r'[\W_]+', '', s1).lower()
    norm2 = re.sub(r'[\W_]+', '', s2).lower()
    return norm1 == norm2


def extract_text_with_location(filepath):
    """
    Extracts text and tracks the page or paragraph number.
    Returns a list of tuples: [(location_string, sentence_text), ...]
    """
    _, ext = os.path.splitext(filepath.lower())
    lines_with_loc = []
    
    def process_block(text_block, loc_label):
        """Helper to safely split by newline first, then by sentence."""
        # Split by newline FIRST to preserve table rows and bullet points from merging
        for line in text_block.split('\n'):
            cleaned_line = sanitize_text(line)
            if cleaned_line:
                # Split sentences if it's a long paragraph. 
                # Table rows usually lack punctuation, so they remain intact as a single row string.
                sentences = re.split(r'(?<=[.!?]) +', cleaned_line)
                for s in sentences:
                    if s.strip():
                        lines_with_loc.append((loc_label, s.strip()))

    if ext == '.pdf':
        with open(filepath, 'rb') as file:
            reader = PyPDF2.PdfReader(file)
            for i, page in enumerate(reader.pages):
                text = page.extract_text()
                if text:
                    process_block(text, f"Page {i+1}")
                    
    elif ext == '.docx':
        doc = docx.Document(filepath)
        # 1. Process standard paragraphs
        for i, para in enumerate(doc.paragraphs):
            process_block(para.text, f"Para {i+1}")
            
        # 2. Process tables (python-docx keeps tables completely separate from paragraphs)
        for i, table in enumerate(doc.tables):
            for j, row in enumerate(table.rows):
                # Join cells with a pipe delimiter for clean tabular readability
                row_text = " | ".join(cell.text.replace('\n', ' ').strip() for cell in row.cells if cell.text)
                process_block(row_text, f"Table {i+1} Row {j+1}")
                
    elif ext == '.txt':
        with open(filepath, 'r', encoding='utf-8') as f:
            for i, line in enumerate(f):
                process_block(line, f"Line {i+1}")
    else:
        raise ValueError(f"Unsupported file format: {ext}. Please use PDF, DOCX, or TXT.")

    return lines_with_loc


def word_level_diff(text1, text2):
    """
    Performs a word-level comparison between two strings to highlight exactly
    what changed within a modified block.
    """
    words1 = text1.split()
    words2 = text2.split()
    matcher = difflib.SequenceMatcher(None, words1, words2)
    
    res1, res2 = [], []
    has_real_change = False
    
    for tag, i1, i2, j1, j2 in matcher.get_opcodes():
        w1_chunk = " ".join(words1[i1:i2])
        w2_chunk = " ".join(words2[j1:j2])
        
        if tag == 'equal':
            res1.append(w1_chunk)
            res2.append(w2_chunk)
        elif tag == 'delete':
            # Ignore deletions that are just random punctuation or spaces
            if is_false_positive(w1_chunk, ""):
                res1.append(w1_chunk)
            else:
                res1.append(f"<del>{w1_chunk}</del>")
                has_real_change = True
        elif tag == 'insert':
            # Ignore insertions that are just random punctuation or spaces
            if is_false_positive("", w2_chunk):
                res2.append(w2_chunk)
            else:
                res2.append(f"<ins>{w2_chunk}</ins>")
                has_real_change = True
        elif tag == 'replace':
            # If the difference is just spacing/capitalization (like "pow er" vs "power"), ignore it
            if is_false_positive(w1_chunk, w2_chunk):
                # Output the clean version without highlight
                res1.append(w1_chunk)
                res2.append(w2_chunk)
            else:
                res1.append(f"<del>{w1_chunk}</del>")
                res2.append(f"<ins>{w2_chunk}</ins>")
                has_real_change = True
                
    return " ".join(res1), " ".join(res2), has_real_change


def generate_diff_report(file1_path, file2_path, output_html="diff_report.html"):
    """
    Compares the texts and generates a custom, visually appealing 3-column HTML table.
    """
    print(f"Extracting text from: {file1_path}")
    data1 = extract_text_with_location(file1_path)
    
    print(f"Extracting text from: {file2_path}")
    data2 = extract_text_with_location(file2_path)
    
    print("Comparing documents...")
    text1 = [item[1] for item in data1]
    text2 = [item[1] for item in data2]
    
    matcher = difflib.SequenceMatcher(None, text1, text2)
    
    html_rows = []
    
    for tag, i1, i2, j1, j2 in matcher.get_opcodes():
        if tag == 'equal':
            continue  # Skip unchanged text
            
        # Determine locations safely
        loc1 = data1[i1][0] if i1 < len(data1) else (data1[-1][0] if data1 else "End of Doc")
        loc2 = data2[j1][0] if j1 < len(data2) else (data2[-1][0] if data2 else "End of Doc")
        
        block1_text = " ".join(text1[i1:i2])
        block2_text = " ".join(text2[j1:j2])
        
        # Word-level highlight for replacements
        if tag == 'replace':
            content1, content2, has_real_change = word_level_diff(block1_text, block2_text)
            if not has_real_change:
                continue # The "difference" was purely formatting/spacing. Skip row entirely!
            status_badge = f"<span class='badge badge-replace'>Modified</span>"
        elif tag == 'delete':
            if is_false_positive(block1_text, ""):
                continue # Skip row if only punctuation/spacing was deleted
            content1 = f"<del>{block1_text}</del>"
            content2 = ""
            status_badge = f"<span class='badge badge-delete'>Deleted</span>"
        elif tag == 'insert':
            if is_false_positive("", block2_text):
                continue # Skip row if only punctuation/spacing was inserted
            content1 = ""
            content2 = f"<ins>{block2_text}</ins>"
            status_badge = f"<span class='badge badge-insert'>Inserted</span>"

        location_cell = f"<strong>Doc 1:</strong> {loc1}<br><strong>Doc 2:</strong> {loc2}<br><br>{status_badge}"
        
        row = f"<tr><td>{location_cell}</td><td>{content1}</td><td>{content2}</td></tr>"
        html_rows.append(row)

    if not html_rows:
        html_rows.append("<tr><td colspan='3' style='text-align: center; padding: 30px;'><strong>No substantial text differences found!</strong><br><small>(Only minor formatting, spacing, or punctuation changes were detected).</small></td></tr>")

    # HTML Boilerplate
    filename1 = os.path.basename(file1_path)
    filename2 = os.path.basename(file2_path)
    
    html_template = f"""
    <!DOCTYPE html>
    <html lang="en">
    <head>
        <meta charset="UTF-8">
        <meta name="viewport" content="width=device-width, initial-scale=1.0">
        <title>Document Comparison Report</title>
        <style>
            body {{ font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, Helvetica, Arial, sans-serif; background-color: #f4f7f9; color: #333; margin: 0; padding: 20px; }}
            h2 {{ text-align: center; color: #2c3e50; }}
            .container {{ max-width: 1200px; margin: 0 auto; background: #fff; padding: 20px; border-radius: 8px; box-shadow: 0 4px 6px rgba(0,0,0,0.1); }}
            table {{ width: 100%; border-collapse: collapse; margin-top: 20px; table-layout: fixed; }}
            th, td {{ border: 1px solid #ddd; padding: 12px; text-align: left; vertical-align: top; word-wrap: break-word; }}
            th {{ background-color: #2c3e50; color: #fff; position: sticky; top: 0; z-index: 10; }}
            th:nth-child(1) {{ width: 15%; }}
            th:nth-child(2), th:nth-child(3) {{ width: 42.5%; }}
            tr:nth-child(even) {{ background-color: #f9f9f9; }}
            del {{ background-color: #ffe6e6; color: #d32f2f; text-decoration: line-through; padding: 2px 4px; border-radius: 3px; font-weight: bold; }}
            ins {{ background-color: #e6ffe6; color: #2e7d32; text-decoration: none; padding: 2px 4px; border-radius: 3px; font-weight: bold; border-bottom: 2px solid #2e7d32; }}
            .badge {{ display: inline-block; padding: 4px 8px; border-radius: 12px; font-size: 0.85em; font-weight: bold; color: #fff; }}
            .badge-replace {{ background-color: #f39c12; }}
            .badge-delete {{ background-color: #e74c3c; }}
            .badge-insert {{ background-color: #2ecc71; }}
        </style>
    </head>
    <body>
        <div class="container">
            <h2>Document Audit Report</h2>
            <table>
                <thead>
                    <tr>
                        <th>Location & Status</th>
                        <th>Source 1: {filename1}</th>
                        <th>Source 2: {filename2}</th>
                    </tr>
                </thead>
                <tbody>
                    {"".join(html_rows)}
                </tbody>
            </table>
        </div>
    </body>
    </html>
    """
    
    with open(output_html, 'w', encoding='utf-8') as f:
        f.write(html_template)
        
    print(f"Comparison complete! Open '{output_html}' in your web browser to view the beautiful tabular differences.")


if __name__ == "__main__":
    if len(sys.argv) == 3:
        # Allow command-line execution for automation if desired
        doc_a = sys.argv[1]
        doc_b = sys.argv[2]
    else:
        # Fall back to the Pop-Up Browser UI
        print("No command-line arguments provided. Launching graphical file picker...")
        
        # Initialize tkinter and hide the main background window
        root = tk.Tk()
        root.withdraw()
        
        print("Please select the FIRST document (Source 1)...")
        doc_a = filedialog.askopenfilename(
            title="Select First Document (Source 1)",
            filetypes=[("Supported Files", "*.pdf *.docx *.txt"), ("All Files", "*.*")]
        )
        
        if not doc_a:
            print("No file selected for Document 1. Exiting.")
            sys.exit(0)
            
        print("Please select the SECOND document (Source 2)...")
        doc_b = filedialog.askopenfilename(
            title="Select Second Document (Source 2)",
            filetypes=[("Supported Files", "*.pdf *.docx *.txt"), ("All Files", "*.*")]
        )
        
        if not doc_b:
            print("No file selected for Document 2. Exiting.")
            sys.exit(0)

    try:
        generate_diff_report(doc_a, doc_b, "comparison_results.html")
    except Exception as e:
        print(f"An error occurred: {e}")